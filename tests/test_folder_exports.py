"""The Folder column in the test-case and checklist exports.

Folders live only in the database, so an export used to drop them: a pack
filed under "Header / Why Us / Careers page" came out flat. Every format now
carries the full path, appended as the last column so a consumer reading
columns by position is unaffected, and an unfiled item gets an empty cell.

The other half of the contract is the way back in: the import matcher must
treat "Folder" as an ignored column, not mistake it for one it knows.
"""

from __future__ import annotations

import csv
import io
import secrets

import pytest
from openpyxl import load_workbook

from engine import db
from engine import imports

_RUN = secrets.token_hex(4)
_PATH = "Header / Why Us / Careers page"


def _case(external_id):
    return {
        "id": external_id, "section": "Auth", "section_num": 1,
        "summary": f"Verify that {external_id} works",
        "preconditions": "", "test_steps": "1. Open the page",
        "test_data": "", "expected_result": "The page opens",
        "issues": "", "comment": "", "user_story_id": "",
        "category": "Positive", "priority": "High", "status": "Unchecked",
        "testing_type": "Functional",
    }


def _item(external_id):
    return {"id": external_id, "section": "Header", "objective":
            f"Verify that {external_id} is correct", "category": "Positive",
            "priority": "High", "status": "Unchecked", "item_num": "",
            "depth": 2}


@pytest.fixture
def project(app, request):
    pid = db.upsert_project(name=f"Folder export {request.node.name} {_RUN}"[:180])
    db.save_test_cases(pid, [_case("TC-001"), _case("TC-002")])
    db.save_checklist(pid, [_item("HDR_001"), _item("HDR_002")])
    header = db.create_folder(pid, "Header")
    why = db.create_folder(pid, "Why Us", header["id"])
    careers = db.create_folder(pid, "Careers page", why["id"])
    db.assign_to_folder(pid, "test_case", ["TC-001"], careers["id"])
    db.assign_to_folder(pid, "checklist_item", ["HDR_001"], careers["id"])
    return pid


@pytest.fixture
def use_project(client, project):
    with client.session_transaction() as sess:
        sess["project_id"] = project
    return client


def _get(client, fmt):
    resp = client.get(f"/export/{fmt}")
    assert resp.status_code == 200, resp.data[:300]
    return resp


def _csv_rows(resp):
    return list(csv.reader(io.StringIO(resp.get_data(as_text=True))))


def _xlsx_rows(resp):
    ws = load_workbook(io.BytesIO(resp.data)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_folder_paths_name_the_whole_chain(project):
    assert db.folder_paths(project, "test_case") == {"TC-001": _PATH}
    assert db.folder_paths(project, "checklist_item") == {"HDR_001": _PATH}


@pytest.mark.parametrize("fmt,id_col", [
    ("csv-testcases", 0), ("csv-checklist", 1)])
def test_csv_appends_folder_last(use_project, fmt, id_col):
    rows = _csv_rows(_get(use_project, fmt))
    assert rows[0][-1] == "Folder"
    # The existing columns are where they were.
    assert rows[0][0] == ("TC ID" if fmt == "csv-testcases" else "No")
    by_id = {r[id_col]: r[-1] for r in rows[1:]}
    filed, unfiled = (("TC-001", "TC-002") if fmt == "csv-testcases"
                      else ("HDR_001", "HDR_002"))
    assert by_id[filed] == _PATH
    assert by_id[unfiled] == ""


@pytest.mark.parametrize("fmt,id_col", [
    ("xlsx-testcases", 0), ("xlsx-checklist", 1)])
def test_xlsx_appends_folder_last(use_project, fmt, id_col):
    rows = _xlsx_rows(_get(use_project, fmt))
    assert rows[0][-1] == "Folder"
    by_id = {r[id_col]: r[-1] for r in rows[1:]}
    filed, unfiled = (("TC-001", "TC-002") if fmt == "xlsx-testcases"
                      else ("HDR_001", "HDR_002"))
    assert by_id[filed] == _PATH
    assert by_id[unfiled] in ("", None)


def test_markdown_shows_the_folder(use_project):
    text = _get(use_project, "markdown").get_data(as_text=True)
    tc1 = text.split("#### TC-001")[1].split("####")[0]
    tc2 = text.split("#### TC-002")[1].split("---")[0]
    assert f"- **Folder:** {_PATH}" in tc1
    assert "Folder" not in tc2
    assert "| Comments | Folder |" in text
    row = next(ln for ln in text.splitlines() if "| HDR_001 |" in ln)
    assert row.rstrip().endswith(f"| {_PATH} |")


def test_html_shows_the_folder(use_project):
    text = _get(use_project, "html").get_data(as_text=True)
    assert f"Folder: {_PATH}" in text
    assert "<th>Comments</th><th>Folder</th>" in text
    row = next(ln for ln in text.splitlines() if "<td>HDR_001</td>" in ln)
    assert f"<td>{_PATH}</td></tr>" in row


@pytest.mark.parametrize("fmt,ext,parse", [
    ("csv-testcases", ".csv", imports.parse_test_cases),
    ("xlsx-testcases", ".xlsx", imports.parse_test_cases),
    ("csv-checklist", ".csv", imports.parse_checklist),
    ("xlsx-checklist", ".xlsx", imports.parse_checklist),
    ("markdown", ".md", imports.parse_test_cases),
])
def test_an_export_still_imports(use_project, tmp_path, fmt, ext, parse):
    """The Folder column is ignored on the way back in, not mis-mapped."""
    path = tmp_path / f"pack{ext}"
    path.write_bytes(_get(use_project, fmt).data)
    rows = parse(str(path), path.name)
    ids = {r.id for r in rows}
    if "checklist" in fmt:
        assert {"HDR_001", "HDR_002"} <= ids
    else:
        assert {"TC-001", "TC-002"} <= ids
    for r in rows:
        for value in vars(r).values():
            assert value != _PATH, (fmt, r)


@pytest.mark.parametrize("kind", ["test_cases", "checklist"])
def test_the_import_preview_ignores_folder(kind):
    from engine import import_preview
    header = (["TC ID", "Section", "Summary", "Status", "Folder"]
              if kind == "test_cases"
              else ["No", "ID", "Section", "Objective", "Status", "Folder"])
    analysis = import_preview.analyse(kind, header)
    assert "Folder" in analysis.ignored
